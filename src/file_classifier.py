"""Content-first file classification helpers.

The classifier evaluates declarative profiles against workbook content.
Each profile may declare sheet signals, required columns, and markers.
"""

from __future__ import annotations

from pathlib import Path
from typing import Any


def classify_file(file_path: str | Path, profiles: list[dict[str, Any]]) -> dict[str, Any] | None:
    """Classify an Excel workbook using content-first profile matching.

    Returns the profile result dict when a profile matches, else None.
    """
    path = Path(file_path)
    if path.suffix.lower() not in (".xlsx", ".xlsm", ".xls"):
        return None

    try:
        from openpyxl import load_workbook

        wb = load_workbook(str(path), read_only=True, data_only=True)
    except Exception:
        return None

    try:
        sheet_names = list(wb.sheetnames)
        sheet_name_map = {s.upper(): s for s in sheet_names}

        for profile in profiles:
            matched_sheets = _find_matching_sheets(
                sheet_name_map, profile.get("sheet_signals", []),
            )
            if matched_sheets is None:
                continue

            required_columns = profile.get("required_columns", [])
            if required_columns and not _check_columns(wb, matched_sheets, required_columns):
                continue

            marker = profile.get("content_marker")
            if marker and not _check_content(wb, matched_sheets, marker):
                continue

            result = profile.get("result")
            if isinstance(result, dict):
                return result
        return None
    finally:
        wb.close()


def _find_matching_sheets(
    sheet_name_map: dict[str, str], signals: list[str],
) -> list[str] | None:
    """Return matched sheet names or None when profile signals do not match."""
    if not signals:
        return list(sheet_name_map.values())

    matched: list[str] = []
    for signal in signals:
        signal_upper = str(signal).upper()
        hits = [
            original for upper, original in sheet_name_map.items() if signal_upper in upper
        ]
        if not hits:
            return None
        matched.extend(hits)

    # Preserve order while deduplicating.
    seen: set[str] = set()
    unique: list[str] = []
    for name in matched:
        if name not in seen:
            seen.add(name)
            unique.append(name)
    return unique


def _normalise_header_name(value: Any) -> str:
    return str(value).strip().lower().replace(" ", "_")


def _check_columns(wb, matched_sheets: list[str], required: list[str]) -> bool:
    """Validate that any matched sheet contains all required headers."""
    required_norm = {_normalise_header_name(column) for column in required}
    for sheet_name in matched_sheets:
        ws = wb[sheet_name]
        header = next(ws.iter_rows(min_row=1, max_row=1, values_only=True), None)
        if not header:
            continue
        header_norm = {_normalise_header_name(v) for v in header if v is not None}
        if required_norm.issubset(header_norm):
            return True
    return False


def _check_content(wb, matched_sheets: list[str], marker: dict[str, Any]) -> bool:
    """Validate that marker content exists in any matched sheet."""
    for sheet_name in matched_sheets:
        ws = wb[sheet_name]

        if "column" in marker and "contains" in marker:
            header = next(ws.iter_rows(min_row=1, max_row=1, values_only=True), None)
            if not header:
                continue
            target_col = _normalise_header_name(marker["column"])
            col_idx = None
            for idx, value in enumerate(header):
                if value is not None and _normalise_header_name(value) == target_col:
                    col_idx = idx
                    break
            if col_idx is not None:
                needle = str(marker["contains"]).upper()
                for row in ws.iter_rows(min_row=2, max_row=50, values_only=True):
                    cell = str(row[col_idx] or "").upper()
                    if needle in cell:
                        return True

        if "any_cell_contains" in marker:
            needle = str(marker["any_cell_contains"]).upper()
            max_rows = int(marker.get("max_rows", 20))
            for row in ws.iter_rows(min_row=1, max_row=max_rows, values_only=True):
                for cell in row:
                    if cell is not None and needle in str(cell).upper():
                        return True

    return False
